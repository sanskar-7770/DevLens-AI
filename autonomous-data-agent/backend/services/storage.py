import csv
import io
import json
import re
from pathlib import Path
from uuid import uuid4, UUID
from zipfile import ZipFile
import numpy as np
import pandas as pd
from fastapi import HTTPException
from openpyxl import load_workbook
from .analysis import analyze

UPLOADS = Path(__file__).resolve().parents[1] / "uploads"
MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS, MAX_COLUMNS, MAX_CELLS = 100_000, 200, 2_000_000

def validate_headers(headers):
    names = [str(c).strip() if c is not None else "" for c in headers]
    if not all(names) or len(set(names)) != len(names):
        raise ValueError("Column headers must be non-empty and unique.")

def load_dataset(content, extension):
    try:
        if extension == ".csv":
            text = content.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text), strict=True)
            rows = []
            width = None
            for row in reader:
                if not row: continue
                if width is None: width = len(row)
                if len(row) != width: raise ValueError("CSV rows must have the same number of fields as the header.")
                rows.append(row)
                if width > MAX_COLUMNS or len(rows) > MAX_ROWS+1 or (len(rows)-1)*width > MAX_CELLS:
                    raise ValueError("Dataset exceeds the processing limits.")
            if len(rows) < 2: raise ValueError("Dataset must contain a header and at least one data row.")
            validate_headers(rows[0])
            df = pd.read_csv(io.StringIO(text), skip_blank_lines=True)
        else:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100*1024*1024:
                    raise ValueError("Excel workbook expands beyond the 100 MB safety limit.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS+1:
                    raise ValueError("Excel sheet exceeds the processing limits.")
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(row)
                    if len(row) > MAX_COLUMNS or len(rows) > MAX_ROWS+1 or (len(rows)-1)*len(row) > MAX_CELLS:
                        raise ValueError("Dataset exceeds the processing limits.")
                if len(rows) < 2: raise ValueError("Dataset must contain a header and at least one data row.")
                validate_headers(rows[0])
                df = pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]]).replace("", np.nan)
            finally: workbook.close()
        if df.empty or df.dropna(how="all").empty: raise ValueError("Dataset has no non-empty data rows.")
        if df.size > MAX_CELLS: raise ValueError("Dataset exceeds the processing limits.")
        return df.replace([np.inf, -np.inf], np.nan)
    except (ValueError, UnicodeError, csv.Error) as error:
        raise HTTPException(422, str(error)) from error
    except Exception as error:
        raise HTTPException(422, "Unable to read this file. Upload a valid UTF-8 CSV or XLSX workbook.") from error

def save_dataset(content, filename):
    extension = Path(filename).suffix.lower()
    if extension not in (".csv", ".xlsx"): raise HTTPException(415, "Supported file types are CSV and XLSX.")
    if not content: raise HTTPException(422, "The uploaded file is empty.")
    if len(content) > MAX_BYTES: raise HTTPException(413, "Maximum upload size is 25 MB.")
    report = analyze(load_dataset(content, extension))
    dataset_id = str(uuid4())
    safe_name = re.sub(r"[^a-zA-Z0-9_. -]", "_", filename.replace("\\", "/").split("/")[-1])[:180]
    report["overview"].update(dataset_id=dataset_id, filename=safe_name, file_type=extension[1:])
    UPLOADS.mkdir(exist_ok=True)
    target = UPLOADS / f"{dataset_id}.json"
    temporary = target.with_suffix(".tmp")
    source = UPLOADS / f"{dataset_id}{extension}"
    try:
        source.write_bytes(content)
        temporary.write_text(json.dumps(report, allow_nan=False), encoding="utf-8")
        temporary.replace(target)
    except Exception:
        source.unlink(missing_ok=True)
        raise
    finally: temporary.unlink(missing_ok=True)
    return report["overview"]

def get_report(dataset_id):
    try: canonical = str(UUID(dataset_id))
    except ValueError: raise HTTPException(404, "Dataset not found.")
    path = UPLOADS / f"{canonical}.json"
    if not path.is_file(): raise HTTPException(404, "Dataset not found. Upload the file again.")
    return json.loads(path.read_text(encoding="utf-8"))


def get_dataframe(dataset_id):
    """Reload the validated source, never the 20-row preview or a pickle."""
    report = get_report(dataset_id)
    extension = "." + report["overview"]["file_type"]
    if extension not in (".csv", ".xlsx"):
        raise HTTPException(422, "Unsupported stored dataset format.")
    source = UPLOADS / f"{UUID(dataset_id)}{extension}"
    if not source.is_file():
        raise HTTPException(409, "This Phase 1 report has no retained dataset. Re-upload the original file to use Ask Your Data; the existing dashboard still works.")
    return load_dataset(source.read_bytes(), extension)
